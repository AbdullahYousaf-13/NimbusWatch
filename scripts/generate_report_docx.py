from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT_DIR = Path(__file__).resolve().parents[1]
SOURCE_PATH = ROOT_DIR / "docs" / "NIMBUSWATCH_FINAL_REPORT.md"
OUTPUT_PATH = ROOT_DIR / "docs" / "NIMBUSWATCH_FINAL_REPORT.docx"


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def clean_inline(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("`", "")
    return text.strip()


def build_document(lines: list[str]) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for style_name, size in [("Title", 20), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NimbusWatch Final Report")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(20)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Cloud Computing Course Project").italic = True

    document.add_paragraph()

    body_started = False
    in_code_block = False
    code_lines: list[str] = []
    bullet_pattern = re.compile(r"^[-*]\s+(.*)")
    numbered_pattern = re.compile(r"^\d+\.\s+(.*)")

    for raw_line in lines:
        line = raw_line.rstrip()

        if line.startswith("```"):
            if in_code_block:
                paragraph = document.add_paragraph(style="Intense Quote")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not line.strip():
            if body_started:
                document.add_paragraph("")
            continue

        if line.startswith("# "):
            if not body_started:
                body_started = True
                continue
            paragraph = document.add_paragraph(clean_inline(line[2:]), style="Heading 1")
            paragraph.paragraph_format.space_before = Pt(12)
            continue

        if line.startswith("## "):
            document.add_paragraph(clean_inline(line[3:]), style="Heading 1")
            body_started = True
            continue

        if line.startswith("### "):
            document.add_paragraph(clean_inline(line[4:]), style="Heading 2")
            body_started = True
            continue

        if line.startswith("#### "):
            document.add_paragraph(clean_inline(line[5:]), style="Heading 3")
            body_started = True
            continue

        bullet_match = bullet_pattern.match(line)
        if bullet_match:
            document.add_paragraph(clean_inline(bullet_match.group(1)), style="List Bullet")
            body_started = True
            continue

        numbered_match = numbered_pattern.match(line)
        if numbered_match:
            document.add_paragraph(clean_inline(numbered_match.group(1)), style="List Number")
            body_started = True
            continue

        paragraph = document.add_paragraph()
        paragraph.add_run(clean_inline(line))
        body_started = True

    return document


def main() -> None:
    lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
    document = build_document(lines)
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)
    document.save(OUTPUT_PATH)
    print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
